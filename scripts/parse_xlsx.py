"""Parse the workout spreadsheet into src/data/seed.json.

Usage: python3 scripts/parse_xlsx.py "MVP - Michael Olson.xlsx"
Requires openpyxl (pip install openpyxl).

Sheet rules (confirmed with Michael 10/5):
- Excel auto-formatted rep ranges as dates (8-10 -> Aug 10). Convert back to "min-max".
  "10-8" is a typo for 8-10, so ranges are always sorted low-high.
- A row with a slot (1a, 2b) is an exercise. Rows below it with no slot are variations;
  each variation is its own exercise with its own sets (parent row becomes a name prefix).
- Merged Sets/Reps/Weight cells apply to every row in the merge (tri-sets 6a/6b/6c).
- Blank sets default to 3. Blank weight stays blank.
- Column A = exercise tips (editable in app).
- Columns I-L = a logged session (Sets, ?, Reps, Weight) on the day's header date.
- Supersets = same number + different letter. Back day 5c is standalone.
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "src" / "data" / "seed.json"
STANDALONE_SLOTS = {("back-triceps", "5c")}
TYPOS = {"cbale": "cable", "tripl edrop": "triple drop", "amrs": "arms", "abrollout": "ab rollout",
         "assited": "assisted", "aseat": "seat", "wth": "with", "Seat#": "seat #"}


def clean(s):
    for bad, good in TYPOS.items():
        s = re.sub(rf"\b{re.escape(bad)}", good, s)
    return re.sub(r"(\d+)\*", "\\1°", s).strip()

TAG_RULES = [
    (r"curl", "Biceps", "Pull"),
    (r"tricep|skull|triceps", "Triceps", "Push"),
    (r"shrug|rear fly", "Shoulders", "Pull"),
    (r"pullup|pulldown|row", "Back", "Pull"),
    (r"bench|press|fly|dip", "Chest", "Push"),
]


def tag(name):
    n = name.lower()
    for pat, muscle, move in TAG_RULES:
        if re.search(pat, n):
            return muscle, move
    return "Other", "Push"


def reps_of(v):
    """Return (min, max, text) for a reps cell."""
    if v is None:
        return None, None, ""
    if isinstance(v, (dt.datetime, dt.date)):
        lo, hi = sorted([v.month, v.day])
        return lo, hi, f"{lo}-{hi}"
    if isinstance(v, (int, float)):
        return int(v), int(v), str(int(v))
    s = str(v).strip()
    nums = [int(x) for x in re.findall(r"\d+", s)]
    if not nums:
        return None, None, s
    return min(nums), max(nums), s


def weight_of(v):
    """Return weight dict: value (first number), perSide, drops, text."""
    if v is None:
        return {"value": None, "perSide": False, "drops": None, "text": ""}
    if isinstance(v, (int, float)):
        return {"value": float(v), "perSide": False, "drops": None, "text": f"{v:g}"}
    s = str(v).strip()
    per_side = bool(re.search(r"\d\s*ea\b", s))
    nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", s)]
    drops = nums if "," in s and len(nums) > 1 else None
    text = re.sub(r"(\d)\s*ea\b", r"\1 ea", s)
    return {"value": nums[0] if nums else None, "perSide": per_side, "drops": drops, "text": text}


def merged_lookup(ws):
    """Map every cell coordinate inside a merge to the merge's top-left value."""
    m = {}
    for rng in ws.merged_cells.ranges:
        tl = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = tl
    return m


def slugify(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def main(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    merged = merged_lookup(ws)

    def val(r, c):
        v = ws.cell(r, c).value
        return merged.get((r, c), v) if v is None else v

    exercises, sessions = [], []
    day = None
    parent = None  # last slotted row, for variations
    var_idx = 0
    history = {}  # day id -> list of exercise logs

    for r in range(1, ws.max_row + 1):
        a, c, d = ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        # Day header: column C title + date in E, column A says "Michael Notes"
        if a == "Michael Notes" and c:
            label = str(c).strip().replace("Ticeps", "Triceps")
            date = val(r, 5)
            day = {"id": slugify(label.replace("and", "&").replace("&", "")), "label": label,
                   "date": date.date().isoformat() if isinstance(date, dt.datetime) else None}
            day["id"] = "chest-bicep" if "chest" in label.lower() else "back-triceps"
            history[day["id"]] = {"day": day, "logs": []}
            parent = None
            continue
        if day is None or d is None or str(d).strip() in ("Exercise",):
            continue
        name = clean(str(d))
        slot = str(c).strip() if c else None
        tip = clean(str(a)) if a else ""

        # Variations only hang off a parent row that has no sets/reps/weight of its own;
        # anything else without a slot is a section label like "cables"
        if slot is None and (parent is None or parent["has_data"]):
            continue

        sets_v, reps_v, wt_v = val(r, 5), val(r, 6), val(r, 7)

        if slot:
            has_data = any(ws.cell(r, cc).value is not None or (r, cc) in merged for cc in (5, 6, 7))
            parent = {"slot": slot, "name": name, "tip": tip, "row": r, "has_data": has_data}
            var_idx = 0
            # Parent with no data that has variations underneath: skip, variations carry it
            nxt = ws.cell(r + 1, 4).value is not None and ws.cell(r + 1, 3).value is None
            if not has_data and nxt:
                continue
            ex_name, ex_slot, parent_name = name, slot, None
        else:
            var_idx += 1
            ex_name = f"{parent['name']} ({name})"
            ex_slot = f"{parent['slot']}.{var_idx}"
            parent_name = parent["name"]
            tip = " ".join(t for t in [parent["tip"], tip] if t)

        group_num = int(re.match(r"\d+", ex_slot).group())
        letter = re.match(r"\d+([a-z])", ex_slot).group(1)
        if (day["id"], parent["slot"]) in STANDALONE_SLOTS or parent_name and _parent_alone(ws, parent["row"]):
            group = f"{day['id']}-{ex_slot}"
        elif parent_name:
            group = f"{day['id']}-{group_num}-v{var_idx}"  # pair variation i of 3a with variation i of 3b
        else:
            group = f"{day['id']}-{group_num}"

        lo, hi, rtext = reps_of(reps_v)
        muscle, move = tag(ex_name)
        ex_id = f"{day['id']}-{ex_slot}"
        exercises.append({
            "id": ex_id,
            "day": day["id"],
            "dayLabel": day["label"],
            "slot": ex_slot,
            "letter": letter,
            "group": group,
            "order": len(exercises),
            "name": ex_name,
            "parent": parent_name,
            "sets": int(sets_v) if isinstance(sets_v, (int, float)) else 3,
            "repsMin": lo,
            "repsMax": hi,
            "repsText": rtext,
            "weight": weight_of(wt_v),
            "muscle": muscle,
            "movement": move,
            "tips": tip,
        })

        # Columns I-L: logged session
        i_sets, j, k, l = (ws.cell(r, cc).value for cc in (9, 10, 11, 12))
        if i_sets is not None:
            reps_cell = k
            wt_cell = l
            if isinstance(j, (dt.datetime, dt.date)):  # shifted row: J=reps range, K=weight
                reps_cell, wt_cell = j, k
            lo2, hi2, _ = reps_of(reps_cell)
            reps = round((lo2 + hi2) / 2) if lo2 is not None else None
            w = weight_of(wt_cell)
            if w["value"] is None:
                w = exercises[-1]["weight"]
            per_side = exercises[-1]["weight"]["perSide"]
            sets = [{"reps": reps or 0, "weight": w["value"], "drops": w["drops"]} for _ in range(int(i_sets))]
            history[day["id"]]["logs"].append({
                "exerciseId": ex_id, "name": ex_name, "perSide": per_side, "sets": sets, "durationSec": None,
            })

    for h in history.values():
        if h["logs"] and h["day"]["date"]:
            sessions.append({
                "id": f"seed-{h['day']['id']}",
                "date": h["day"]["date"],
                "title": h["day"]["label"],
                "startedAt": None,
                "durationSec": None,
                "exercises": h["logs"],
                "imported": True,
            })

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"exercises": exercises, "sessions": sessions}, indent=2))
    print(f"wrote {len(exercises)} exercises, {len(sessions)} sessions -> {OUT}")


def _parent_alone(ws, parent_row):
    """True if this parent's slot has no partner letter in the same group (e.g. 5a seated row)."""
    slot = str(ws.cell(parent_row, 3).value)
    num, letter = re.match(r"(\d+)([a-z])", slot).groups()
    for r in range(max(1, parent_row - 12), parent_row + 12):
        s = ws.cell(r, 3).value
        if s and str(s) != slot and re.match(rf"{num}[a-z]$", str(s)):
            # Partner exists, but 5c is forced standalone, which also leaves 5a alone
            day_row = max(rr for rr in range(1, parent_row + 1) if ws.cell(rr, 1).value == "Michael Notes")
            day_id = "chest-bicep" if "chest" in str(ws.cell(day_row, 3).value).lower() else "back-triceps"
            if (day_id, str(s)) in STANDALONE_SLOTS:
                continue
            if day_row == max(rr for rr in range(1, r + 1) if ws.cell(rr, 1).value == "Michael Notes"):
                return False
    return True


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "MVP - Michael Olson.xlsx")
